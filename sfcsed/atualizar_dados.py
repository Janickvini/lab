import openpyxl
import json
import os
import sys
from datetime import datetime

base_dir = os.path.dirname(os.path.abspath(__file__))

# 1. Localizar arquivo Excel
excel_files = [f for f in os.listdir(base_dir) if f.endswith('.xlsx') and not f.startswith('~$')]
if not excel_files:
    print("Erro: Nenhum arquivo .xlsx encontrado na pasta sfcsed.")
    sys.exit(1)

# Prioriza "SF CSED.xlsx" se existir
if "SF CSED.xlsx" in excel_files:
    target_excel = "SF CSED.xlsx"
else:
    target_excel = sorted(excel_files)[-1]

wb_path = os.path.join(base_dir, target_excel)
print(f"Lendo dados de: {wb_path}...")

try:
    wb = openpyxl.load_workbook(wb_path, data_only=True)
except Exception as e:
    print(f"Erro ao abrir arquivo Excel: {e}")
    sys.exit(1)

# Validação e localização dinâmica das abas
def find_sheet(wb, candidates):
    sheetnames_lower = {s.lower(): s for s in wb.sheetnames}
    for cand in candidates:
        if cand.lower() in sheetnames_lower:
            return wb[sheetnames_lower[cand.lower()]]
    return None

ws_sf = find_sheet(wb, ["Sementes", "SF (Scopus)", "SF"])
ws_scopus = find_sheet(wb, ["Referências", "Referencias", "Scopus 3 (Limpo)", "Scopus 3", "Scopus"])

if ws_sf is None:
    print(f"Erro: Aba de Sementes ('Sementes' ou 'SF (Scopus)') não encontrada na planilha. Abas disponíveis: {wb.sheetnames}")
    sys.exit(1)

if ws_scopus is None:
    print(f"Erro: Aba de Referências ('Referências' ou 'Scopus 3 (Limpo)') não encontrada na planilha. Abas disponíveis: {wb.sheetnames}")
    sys.exit(1)

print(f"Aba de Sementes localizada: '{ws_sf.title}'")
print(f"Aba de Referências localizada: '{ws_scopus.title}'")

def clean_value(val):
    if val is None:
        return ""
    if isinstance(val, (int, float)):
        return val
    s = str(val).strip()
    return s

def get_field(item, candidates, default=""):
    for cand in candidates:
        if cand in item and item[cand] is not None and str(item[cand]).strip() != "":
            return item[cand]
    item_lower = {k.strip().lower(): v for k, v in item.items() if k is not None}
    for cand in candidates:
        k_low = cand.strip().lower()
        if k_low in item_lower and item_lower[k_low] is not None and str(item_lower[k_low]).strip() != "":
            return item_lower[k_low]
    return default

# 1. Processar Artigos / Referências
scopus_rows = list(ws_scopus.iter_rows(values_only=True))
sc_headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(scopus_rows[0])]

articles_dict = {}
articles_list = []

for row in scopus_rows[1:]:
    if not any(row):
        continue
    item = {}
    for idx, val in enumerate(row):
        if idx < len(sc_headers):
            item[sc_headers[idx]] = clean_value(val)
    
    art_id = get_field(item, ["ID", "ID Scopus"])
    if art_id != "" and art_id is not None:
        try:
            art_id_int = int(art_id)
        except:
            art_id_int = art_id
        
        doi_val = get_field(item, ["DOI", "Doi"], "")
        doi_link = get_field(item, ["DOILINK", "DOI Link", "Link"], "")
        if not doi_link and doi_val:
            doi_link = f"https://doi.org/{doi_val}" if not str(doi_val).startswith("http") else str(doi_val)
            
        clean_art = {
            "id": art_id_int,
            "fonte": get_field(item, ["Fonte", "Origem"], ""),
            "authors": get_field(item, ["Authors", "Autores"], "Autores não informados"),
            "title": get_field(item, ["Title", "Título", "Titulo", "Título do Artigo"], "Título não informado"),
            "year": get_field(item, ["Year", "Ano"], ""),
            "sourceTitle": get_field(item, ["Source title", "Periódico / Fonte", "Periodico / Fonte", "Source"], ""),
            "doi": doi_val,
            "doiLink": doi_link,
            "abstract": get_field(item, ["Abstract", "Resumo", "Resumo (Abstract)"], ""),
            "authorKeywords": get_field(item, ["Author Keywords", "Palavras-chave do Autor", "Keywords"], ""),
            "indexKeywords": get_field(item, ["Index Keywords", "Palavras-chave Indexadas"], ""),
            "documentType": get_field(item, ["Document Type", "Tipo de Documento"], ""),
            "associatedSFs": []
        }
        articles_dict[art_id_int] = clean_art
        articles_list.append(clean_art)

# 2. Processar Sementes de Futuro
sf_rows = list(ws_sf.iter_rows(values_only=True))
sf_headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(sf_rows[0])]

sementes_list = []

for row in sf_rows[1:]:
    if not any(row):
        continue
    item = {}
    for idx, val in enumerate(row):
        if idx < len(sf_headers):
            item[sf_headers[idx]] = clean_value(val)
    
    sf_id = get_field(item, ["ID da SF", "ID SF", "ID"], "")
    if not sf_id:
        continue

    # Extrair IDs numéricos de artigos associados
    raw_art_ids = str(get_field(item, ["Artigos Scopus Associados (IDs)", "Artigos Associados (IDs)", "Artigos Associados", "Artigos Scopus", "Artigos (IDs)"], ""))
    parsed_art_ids = []
    for p in raw_art_ids.replace(";", ",").split(","):
        p = p.strip()
        if p.isdigit():
            art_num = int(p)
            parsed_art_ids.append(art_num)
            # Vincular a SF de volta ao artigo
            if art_num in articles_dict and sf_id not in articles_dict[art_num]["associatedSFs"]:
                articles_dict[art_num]["associatedSFs"].append(sf_id)

    # Processar PESTEL (lista limpa e normalizada de dimensões)
    pestel_map = {
        "politico": "Político", "político": "Político",
        "economico": "Econômico", "econômico": "Econômico",
        "social": "Social",
        "tecnologico": "Tecnológico", "tecnológico": "Tecnológico",
        "ambiental": "Ambiental",
        "legal": "Legal", "lega": "Legal"
    }
    raw_pestel = str(get_field(item, ["PESTEL", "Dimensões PESTEL", "Dimensoes PESTEL", "Dimensões"], ""))
    pestel_list = [pestel_map.get(dim.strip().lower(), dim.strip()) for dim in raw_pestel.replace(";", ",").split(",") if dim.strip()]

    # Tratar campo revisada / status
    import re
    raw_revisada = get_field(item, ["Revisada", "Status Revisão", "Revisado", "Status"], 0)
    raw_revisada_str = str(raw_revisada).strip()
    status_lower = raw_revisada_str.lower()

    # Tratar campo Vínculo (NOVA COLUNA DEDICADA)
    raw_vinculo = str(get_field(item, ["Vínculo", "Vinculo", "Vínculo SF", "Vinculo SF", "SFs Relacionadas", "SF Relacionada", "Vínculos"], "")).strip()
    vinculo_list = []
    if raw_vinculo and raw_vinculo.upper() not in ["N/A", "NAN", "NONE", "0", ""]:
        vinculo_list = [v.strip() for v in re.split(r"[,;]+", raw_vinculo) if v.strip()]

    # Tratar campo justificativa
    raw_justificativa = str(get_field(item, ["Justificativa", "Justificativa Metodológica", "Motivo"], "")).strip()
    justificativa_val = "" if raw_justificativa.upper() in ["N/A", "NAN", "NONE", ""] else raw_justificativa

    status_tipo = "pendente"
    fusao_id = ""
    divisao_ids = []
    is_revisada = False

    if any(k in status_lower for k in ["divid", "desmembr", "separ", "divis"]):
        status_tipo = "dividida"
        is_revisada = True
        if vinculo_list:
            divisao_ids = [d for d in vinculo_list if d.upper() != sf_id.upper()]
        else:
            m = re.search(r"\((.*?)\)", raw_revisada_str)
            target_text = m.group(1) if m else raw_revisada_str
            extracted = re.findall(r"(SF[A-Za-z0-9]+)", target_text, re.IGNORECASE)
            divisao_ids = [d for d in extracted if d.upper() != sf_id.upper()]
    elif any(k in status_lower for k in ["mescl", "fus", "fundid"]):
        status_tipo = "mesclada"
        is_revisada = True
        if vinculo_list:
            fusao_id = vinculo_list[0]
        else:
            m = re.search(r"\((.*?)\)", raw_revisada_str)
            if m:
                fusao_id = m.group(1).strip()
            else:
                m2 = re.search(r"(SF[A-Z0-9]+)", raw_revisada_str, re.IGNORECASE)
                if m2:
                    fusao_id = m2.group(1).strip()
    elif "descart" in status_lower:
        status_tipo = "descartada"
        is_revisada = True
    elif status_lower in ["ok", "1", "sim", "true", "s", "revisada", "revisado", "aprovada", "aprovado"] or "ok" in status_lower:
        status_tipo = "ok"
        is_revisada = True
    else:
        status_tipo = "pendente"
        is_revisada = False

    if status_tipo == "ok":
        status_label = "Revisada (OK)"
    elif status_tipo == "mesclada":
        status_label = f"Revisada (Mesclada: {fusao_id})" if fusao_id else "Revisada (Mesclada)"
    elif status_tipo == "descartada":
        status_label = "Revisada (Descartada)"
    elif status_tipo == "dividida":
        status_label = f"Revisada (Dividida: {', '.join(divisao_ids)})" if divisao_ids else "Revisada (Dividida)"
    else:
        status_label = "Em revisão"

    # Tratar campo indicadores
    raw_indicadores = str(get_field(item, ["Indicadores", "Indicadores Prospectivos"], "")).strip()
    indicadores_val = "" if raw_indicadores.upper() in ["N/A", "NAN", "NONE", ""] else raw_indicadores

    # Qtd artigos
    qtd_art_raw = get_field(item, ["Qtd Artigos", "Quantidade de Artigos"], 0)
    try:
        qtd_art_num = int(qtd_art_raw)
    except:
        qtd_art_num = 0

    sf_obj = {
        "id": sf_id,
        "fonte": get_field(item, ["Fonte", "Origem"], ""),
        "topico": get_field(item, ["Tópico", "Tópico Mais Amplo Padronizado", "Topico"], "Geral"),
        "titulo": get_field(item, ["Título SF", "Título da Semente de Futuro", "Titulo SF", "Título", "Titulo"], "Sem título"),
        "descricao": get_field(item, ["Descrição (Ref N)", "Descrição Global do Fenômeno (com Citações Numéricas)", "Descricao (Ref N)", "Descrição", "Descricao"], ""),
        "descricaoApa": get_field(item, ["Descrição (Ref APA)", "Descricao (Ref APA)"], ""),
        "classificacao": get_field(item, ["Classificação", "Classificação Prospectiva", "Classificacao"], "Não classificada"),
        "pestel": pestel_list,
        "eixo": get_field(item, ["Eixo", "Eixo CSED", "Eixo Temático"], "Não categorizado"),
        "qtdArtigos": len(parsed_art_ids) if parsed_art_ids else qtd_art_num,
        "artigosIds": parsed_art_ids,
        "referenciasResumidas": get_field(item, ["Referências", "Referências Bibliográficas Resumidas (Autores, Ano, Título)", "Referencias"], ""),
        "implicacoes": get_field(item, ["Implicações", "Conexão / Implicações Estratégicas para o Ensino Superior Privado até 2040", "Implicacoes"], ""),
        "indicadores": indicadores_val,
        "revisada": is_revisada,
        "statusTipo": status_tipo,
        "statusRevisao": raw_revisada_str if raw_revisada_str not in ["0", "None", ""] else "Em revisão",
        "statusLabel": status_label,
        "justificativa": justificativa_val,
        "fusaoId": fusao_id,
        "divisaoIds": divisao_ids,
        "origemDivisao": "",
        "vinculo": vinculo_list,
        "incorporadas": []
    }
    sementes_list.append(sf_obj)

# Mapeamento reverso de incorporações (ex: SFS02 incorpora SF22) e linhagem de divisões (ex: SF31a/b originam de SF31)
sf_map = {s["id"].upper(): s for s in sementes_list}
for s in sementes_list:
    if s["fusaoId"] and s["fusaoId"].upper() in sf_map:
        target = sf_map[s["fusaoId"].upper()]
        if s["id"] not in target["incorporadas"]:
            target["incorporadas"].append(s["id"])
    if s["divisaoIds"]:
        s["divisaoIds"] = [sf_map[d.upper()]["id"] if d.upper() in sf_map else d for d in s["divisaoIds"]]
        for child_id in s["divisaoIds"]:
            if child_id.upper() in sf_map:
                child = sf_map[child_id.upper()]
                child["origemDivisao"] = s["id"]
    if s["vinculo"]:
        for v in s["vinculo"]:
            v_up = v.upper()
            if v_up in sf_map:
                related = sf_map[v_up]
                if related["statusTipo"] == "dividida" and not s["origemDivisao"]:
                    s["origemDivisao"] = related["id"]
                    if s["id"] not in related["divisaoIds"]:
                        related["divisaoIds"].append(s["id"])
                elif related["statusTipo"] == "mesclada" and related["id"] not in s["incorporadas"]:
                    s["incorporadas"].append(related["id"])

# Ordenar SFs numericamente (ex: SF01, SF02... SF31, SF31a, SF31b... SF62, SFQ, SFS)
def sf_sort_key(sf):
    code = sf["id"]
    digits = "".join([c for c in code if c.isdigit()])
    d_int = int(digits) if digits else 9999
    cat = 0 if code.startswith("SF") and not code.startswith("SFQ") and not code.startswith("SFS") else (1 if code.startswith("SFQ") else 2)
    return (cat, d_int, code)

sementes_list.sort(key=sf_sort_key)

# Consolidar dados globais e estatísticas
data_payload = {
    "metadata": {
        "sourceFile": target_excel,
        "totalSF": len(sementes_list),
        "totalArticles": len(articles_list),
        "totalArticlesCited": len([a for a in articles_list if len(a["associatedSFs"]) > 0]),
        "totalRevisadas": len([sf for sf in sementes_list if sf["revisada"]]),
        "totalAprovadas": len([sf for sf in sementes_list if sf["statusTipo"] == "ok"]),
        "totalMescladas": len([sf for sf in sementes_list if sf["statusTipo"] == "mesclada"]),
        "totalDescartadas": len([sf for sf in sementes_list if sf["statusTipo"] == "descartada"]),
        "totalDivididas": len([sf for sf in sementes_list if sf["statusTipo"] == "dividida"]),
        "totalPendentes": len([sf for sf in sementes_list if sf["statusTipo"] == "pendente"]),
        "generatedAt": datetime.now().strftime("%d/%m/%Y %H:%M:%S")
    },
    "sementes": sementes_list,
    "artigos": articles_list
}

# Salvar como data.js
data_js_path = os.path.join(base_dir, "data.js")
with open(data_js_path, "w", encoding="utf-8") as f:
    f.write("// Dados exportados automaticamente de " + target_excel + "\n")
    f.write("const sfData = ")
    json.dump(data_payload, f, ensure_ascii=False, indent=2)
    f.write(";\n")

print(f"\n--- Sincronização Concluída com Sucesso ---")
print(f"Arquivo lido: {target_excel}")
print(f"Sementes de Futuro extraídas: {len(sementes_list)} (Revisadas: {data_payload['metadata']['totalRevisadas']})")
print(f"Artigos / Referências mapeadas: {len(articles_list)} (Citadas em SFs: {data_payload['metadata']['totalArticlesCited']})")
print(f"Arquivo gerado: {data_js_path}")
